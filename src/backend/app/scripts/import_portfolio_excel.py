"""Impor data portofolio dari Excel "Intern Business Strategic - Dashboard Dummy" (Sheet1) ke database.

    python -m app.scripts.import_portfolio_excel path/ke/file.xlsx            # impor
    python -m app.scripts.import_portfolio_excel path/ke/file.xlsx --dry-run  # cek saja, tidak menyimpan

Aturan:
- Satu baris Excel = satu portofolio. Kuncinya "Project ID" (project_code): kalau sudah ada, di-update
  (aman dijalankan berulang); kalau belum, dibuat.
- Customer dicocokkan lewat nama (tanpa beda huruf besar/kecil), dibuat bila belum ada.
- Master data (Customer Size, Industry, Business Unit, Status, Product, Tech Stack) HARUS sudah ada
  lewat seed/halaman master; nilai yang tidak dikenal dilaporkan per baris, bukan dibuat diam-diam.
  Partner tidak seperti itu: partner baru dibuat otomatis.
- Kolom formula (value after tax, progress, total/remaining days, net profit) diabaikan karena dihitung API.
- Bila ada baris bermasalah, tidak ada yang disimpan kecuali dengan --skip-invalid.
"""

import argparse
import asyncio
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.portfolios.calculations import ACTIVE_PROJECT_STATUS_CODE  # noqa: F401  (dokumentasi kode status)
from app.connections.session import AsyncSessionLocal
from app.models.business_unit_model import BusinessUnit
from app.models.company_size_model import CompanySize
from app.models.customer_model import Customer
from app.models.customer_type_model import CustomerType
from app.models.industry_model import Industry
from app.models.legal_structure_model import LegalStructure
from app.models.partner_model import Partner
from app.models.product_model import Product
from app.models.project_financial_model import ProjectFinancial
from app.models.project_model import Project
from app.models.project_partner_model import ProjectPartner
from app.models.project_product_model import ProjectProduct
from app.models.project_status_model import ProjectStatus
from app.models.project_tech_stack_model import ProjectTechStack
from app.models.tech_stack_model import TechStack

HEADER_ROW = 2
FIRST_DATA_ROW = 3

HEADERS = {
    "customer": "Customer",
    "company_size": "Customer Size",
    "annual_revenue": "Revenue/Annual Sales/Annual Budget",
    "customer_type": "Customer Type",
    "industry": "Industry",
    "location": "Location",
    "legal_structure": "Legal Structure",
    "opportunity": "Opportunity",
    "project_name": "Project Name",
    "year": "Year",
    "contract_po_name": "Contract/PO Name",
    "contract_po_number": "Contract/PO Number",
    "business_unit": "Business Unit",
    "project_code": "Project ID",
    "value_before_tax": "Project Value (before tax)",
    "carry_over": "Carry Over Value",
    "status": "Project Status",
    "project_manager": "Project Manager",
    "account_manager": "Account Manager",
    "start_date": "Start Date",
    "end_date": "End Date",
    "objective": "Project Objective",
    "revenue": "Revenue",
    "cogs": "COGS",
    "opex": "Operational Costs",
    "products": "Product Use",
    "frontend": "Frontend",
    "backend": "Backend",
    "data_storage": "Data & Storage",
    "partners": "Partner",
}


# ───────────────────────── parsing (fungsi murni) ─────────────────────────
def normalize(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def clean_text(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"[ \t]+", " ", str(value)).strip()
    return text or None


def split_multi(value) -> list[str]:
    """Pisah daftar dipisah koma/baris baru. Koma di dalam tanda kutip atau kurung tidak memisah,
    contoh: '"Vue.js (Vue, Vue2, Vue3)", Nuxt 3' -> ['Vue.js (Vue, Vue2, Vue3)', 'Nuxt 3']."""
    if value is None:
        return []

    items, buffer, depth, quoted = [], [], 0, False

    for char in str(value):
        if char == '"':
            quoted = not quoted
            continue
        if not quoted:
            if char == "(":
                depth += 1
            elif char == ")":
                depth = max(depth - 1, 0)
            elif char in ",\n;" and depth == 0:
                items.append("".join(buffer))
                buffer = []
                continue
        buffer.append(char)

    items.append("".join(buffer))

    seen, result = set(), []
    for item in items:
        item = re.sub(r"\s+", " ", item).strip()
        if item and normalize(item) not in seen:
            seen.add(normalize(item))
            result.append(item)
    return result


def parse_partners(value) -> list[tuple[str, str]]:
    """'Principal: Zabbix SIA, \\nVendor: PT Bangunindo Teknusa' -> [('Principal', 'Zabbix SIA'), ...]
    Tanpa 'Peran:' dianggap peran 'Partner'."""
    result, seen = [], set()
    for item in split_multi(value):
        role, separator, name = item.partition(":")
        if not separator:
            role, name = "Partner", item
        role, name = role.strip() or "Partner", name.strip()
        if name and normalize(name) not in seen:
            seen.add(normalize(name))
            result.append((role, name))
    return result


def to_decimal(value, field_name: str, errors: list[str], *, required=False) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            errors.append(f"{field_name} kosong")
        return None
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        errors.append(f"{field_name} bukan angka: {value!r}")
        return None
    if number < 0:
        errors.append(f"{field_name} tidak boleh negatif")
        return None
    return number.quantize(Decimal("0.01"))


def to_date(value, field_name: str, errors: list[str]) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    errors.append(f"{field_name} bukan tanggal: {value!r}")
    return None


# ───────────────────────── lookup master data ─────────────────────────
class Masters:
    """Peta nama (normal) -> id untuk tiap master data."""

    def __init__(self):
        self.maps: dict[str, dict[str, object]] = {}
        self.company_size_by_tier: dict[str, object] = {}
        self.tech: dict[tuple[str, str], object] = {}
        self.partners: dict[str, Partner] = {}

    @classmethod
    async def load(cls, db: AsyncSession) -> "Masters":
        self = cls()
        specs = {
            "company_size": (CompanySize, "company_size_name"),
            "customer_type": (CustomerType, "customer_type_name"),
            "industry": (Industry, "industry_name"),
            "legal_structure": (LegalStructure, "legal_structure_name"),
            "business_unit": (BusinessUnit, "business_unit_name"),
            "status": (ProjectStatus, "project_status_name"),
            "product": (Product, "product_name"),
        }
        for key, (model, column) in specs.items():
            rows = (await db.execute(select(model.id, getattr(model, column)))).all()
            self.maps[key] = {normalize(name): item_id for item_id, name in rows}

        for name, item_id in self.maps["company_size"].items():
            tier = re.match(r"tier\s*(\d+)", name)
            if tier:
                self.company_size_by_tier[tier.group(1)] = item_id

        for item_id, category, name in (
            await db.execute(select(TechStack.id, TechStack.tech_stack_category, TechStack.tech_stack_name))
        ).all():
            self.tech[(category, normalize(name))] = item_id

        for partner in (await db.execute(select(Partner))).scalars():
            self.partners[normalize(partner.partner_name)] = partner

        return self

    def find(self, key: str, value: str, label: str, errors: list[str]):
        if key == "company_size":
            tier = re.match(r"tier\s*(\d+)", normalize(value))
            found = self.maps[key].get(normalize(value)) or (
                self.company_size_by_tier.get(tier.group(1)) if tier else None
            )
        else:
            found = self.maps[key].get(normalize(value))
        if found is None:
            errors.append(f"{label} tidak dikenal: {value!r}")
        return found

    def find_tech(self, category: str, names: list[str], errors: list[str]) -> list:
        ids = []
        for name in names:
            found = self.tech.get((category, normalize(name)))
            if found is None:
                errors.append(f"{category} tidak dikenal: {name!r}")
            else:
                ids.append(found)
        return ids


# ───────────────────────── baca Excel ─────────────────────────
@dataclass
class RowResult:
    excel_row: int
    project_code: str | None = None
    errors: list[str] = field(default_factory=list)
    data: dict = field(default_factory=dict)


def read_rows(path: Path) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))

    header = [normalize(cell) for cell in rows[HEADER_ROW - 1]]
    columns, missing = {}, []
    for key, title in HEADERS.items():
        try:
            columns[key] = header.index(normalize(title))
        except ValueError:
            missing.append(title)
    if missing:
        raise SystemExit(f"Kolom tidak ditemukan di baris {HEADER_ROW}: {', '.join(missing)}")

    result = []
    for index, row in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        values = {key: (row[col] if col < len(row) else None) for key, col in columns.items()}
        # Baris kosong (Excel punya ratusan baris berformat tanpa isi) dilewati.
        if not clean_text(values["customer"]) and not clean_text(values["project_code"]):
            continue
        values["_row"] = index
        result.append(values)
    return result


def validate_row(raw: dict, masters: Masters) -> RowResult:
    result = RowResult(excel_row=raw["_row"])
    errors = result.errors

    code = clean_text(raw["project_code"])
    customer_name = clean_text(raw["customer"])
    project_name = clean_text(raw["project_name"])
    result.project_code = code

    if not code:
        errors.append("Project ID kosong")
    elif not re.match(r"^[A-Za-z0-9][A-Za-z0-9_\-./]*$", code):
        errors.append(f"Project ID mengandung karakter tidak valid: {code!r}")
    if not customer_name:
        errors.append("Customer kosong")
    if not project_name:
        errors.append("Project Name kosong")

    year = raw["year"]
    try:
        year = int(float(year))
    except (TypeError, ValueError):
        errors.append(f"Year bukan angka: {year!r}")
        year = None

    value = to_decimal(raw["value_before_tax"], "Project Value (before tax)", errors, required=True)
    carry = to_decimal(raw["carry_over"], "Carry Over Value", errors)
    revenue = to_decimal(raw["revenue"], "Revenue", errors)
    cogs = to_decimal(raw["cogs"], "COGS", errors) or Decimal("0.00")
    opex = to_decimal(raw["opex"], "Operational Costs", errors) or Decimal("0.00")
    start = to_date(raw["start_date"], "Start Date", errors)
    end = to_date(raw["end_date"], "End Date", errors)
    if start and end and end < start:
        errors.append("End Date lebih awal dari Start Date")

    def need(key, value_key, label):
        text = clean_text(raw[value_key])
        if not text:
            errors.append(f"{label} kosong")
            return None
        return masters.find(key, text, label, errors)

    company_size = need("company_size", "company_size", "Customer Size")
    customer_type = need("customer_type", "customer_type", "Customer Type")
    industry = need("industry", "industry", "Industry")
    legal_structure = need("legal_structure", "legal_structure", "Legal Structure")
    business_unit = need("business_unit", "business_unit", "Business Unit")
    status = need("status", "status", "Project Status")

    product_ids = []
    for name in split_multi(raw["products"]):
        found = masters.find("product", name, "Product Use", errors)
        if found is not None:
            product_ids.append(found)

    tech_ids = (
        masters.find_tech("frontend", split_multi(raw["frontend"]), errors)
        + masters.find_tech("backend", split_multi(raw["backend"]), errors)
        + masters.find_tech("data_storage", split_multi(raw["data_storage"]), errors)
    )

    result.data = {
        "customer": {
            "customer_name": customer_name,
            "company_size_id": company_size,
            "annual_revenue": clean_text(raw["annual_revenue"]),
            "customer_type_id": customer_type,
            "industry_id": industry,
            "location": clean_text(raw["location"]),
            "legal_structure_id": legal_structure,
        },
        "project": {
            "project_code": code,
            "project_name": project_name,
            "opportunity": clean_text(raw["opportunity"]),
            "project_year": year,
            "contract_po_name": clean_text(raw["contract_po_name"]),
            "contract_po_number": clean_text(raw["contract_po_number"]),
            "business_unit_id": business_unit,
            "project_status_id": status,
            "value_before_tax": value,
            "carry_over_value": carry,
            "project_manager": clean_text(raw["project_manager"]),
            "account_manager": clean_text(raw["account_manager"]),
            "start_date": start,
            "end_date": end,
            "project_objective": clean_text(raw["objective"]),
        },
        "financial": {
            "revenue": revenue if revenue is not None else value,
            "cogs": cogs,
            "operational_costs": opex,
        },
        "product_ids": list(dict.fromkeys(product_ids)),
        "tech_stack_ids": list(dict.fromkeys(tech_ids)),
        "partners": parse_partners(raw["partners"]),
    }
    return result


# ───────────────────────── simpan ke database ─────────────────────────
@dataclass
class Summary:
    customers_created: int = 0
    projects_created: int = 0
    projects_updated: int = 0
    partners_created: int = 0
    skipped: int = 0


async def apply_row(db: AsyncSession, masters: Masters, result: RowResult, summary: Summary) -> None:
    data = result.data

    name = data["customer"]["customer_name"]
    customer = (
        await db.execute(select(Customer).where(func.lower(Customer.customer_name) == name.lower()))
    ).scalar_one_or_none()
    if customer is None:
        customer = Customer(**data["customer"])
        db.add(customer)
        await db.flush()
        summary.customers_created += 1
    else:
        for key, value in data["customer"].items():
            setattr(customer, key, value)

    project = (
        await db.execute(
            select(Project).where(func.lower(Project.project_code) == result.project_code.lower())
        )
    ).scalar_one_or_none()

    if project is None:
        project = Project(**data["project"], customer_id=customer.id)
        db.add(project)
        await db.flush()
        summary.projects_created += 1
    else:
        for key, value in data["project"].items():
            setattr(project, key, value)
        project.customer_id = customer.id
        summary.projects_updated += 1

    financial = (
        await db.execute(select(ProjectFinancial).where(ProjectFinancial.project_id == project.id))
    ).scalar_one_or_none()
    if financial is None:
        db.add(ProjectFinancial(project_id=project.id, **data["financial"]))
    else:
        for key, value in data["financial"].items():
            setattr(financial, key, value)

    await replace_links(db, project.id, ProjectProduct, "product_id", data["product_ids"])
    await replace_links(db, project.id, ProjectTechStack, "tech_stack_id", data["tech_stack_ids"])

    existing = {
        link.partner_id: link
        for link in (
            await db.execute(select(ProjectPartner).where(ProjectPartner.project_id == project.id))
        ).scalars()
    }
    wanted = {}
    for role, partner_name in data["partners"]:
        partner = masters.partners.get(normalize(partner_name))
        if partner is None:
            partner = Partner(partner_name=partner_name, is_active=True)
            db.add(partner)
            await db.flush()
            masters.partners[normalize(partner_name)] = partner
            summary.partners_created += 1
        wanted[partner.id] = role
    for partner_id, link in existing.items():
        if partner_id not in wanted:
            await db.delete(link)
        else:
            link.partner_role = wanted[partner_id]
    for partner_id, role in wanted.items():
        if partner_id not in existing:
            db.add(ProjectPartner(project_id=project.id, partner_id=partner_id, partner_role=role))

    await db.flush()


async def replace_links(db: AsyncSession, project_id, model, column: str, wanted_ids: list) -> None:
    existing = {
        getattr(link, column): link
        for link in (await db.execute(select(model).where(model.project_id == project_id))).scalars()
    }
    for item_id, link in existing.items():
        if item_id not in wanted_ids:
            await db.delete(link)
    for item_id in wanted_ids:
        if item_id not in existing:
            db.add(model(project_id=project_id, **{column: item_id}))


async def run_import(path: Path, *, dry_run: bool = False, skip_invalid: bool = False) -> tuple[Summary, list[RowResult]]:
    raw_rows = read_rows(path)

    async with AsyncSessionLocal() as db:
        masters = await Masters.load(db)
        results = [validate_row(raw, masters) for raw in raw_rows]

        codes = {}
        for item in results:
            if item.project_code:
                key = item.project_code.lower()
                if key in codes:
                    item.errors.append(f"Project ID dobel dengan baris {codes[key]}")
                else:
                    codes[key] = item.excel_row

        invalid = [item for item in results if item.errors]
        summary = Summary(skipped=len(invalid))

        if invalid and not skip_invalid:
            return summary, results

        for item in results:
            if not item.errors:
                await apply_row(db, masters, item, summary)

        if dry_run:
            await db.rollback()
        else:
            await db.commit()

    return summary, results


def main() -> int:
    parser = argparse.ArgumentParser(description="Impor portofolio dari Excel")
    parser.add_argument("file", type=Path)
    parser.add_argument("--dry-run", action="store_true", help="Validasi dan hitung, tanpa menyimpan")
    parser.add_argument("--skip-invalid", action="store_true", help="Lewati baris bermasalah, simpan sisanya")
    args = parser.parse_args()

    summary, results = asyncio.run(
        run_import(args.file, dry_run=args.dry_run, skip_invalid=args.skip_invalid)
    )

    bad = [item for item in results if item.errors]
    for item in bad:
        print(f"Baris {item.excel_row} ({item.project_code or '-'}):", file=sys.stderr)
        for message in item.errors:
            print(f"  - {message}", file=sys.stderr)

    if bad and not args.skip_invalid:
        print(f"\nDibatalkan: {len(bad)} baris bermasalah, tidak ada yang disimpan.", file=sys.stderr)
        return 1

    mode = "DRY-RUN (tidak disimpan)" if args.dry_run else "Selesai"
    print(
        f"{mode}: customer baru {summary.customers_created}, project baru {summary.projects_created}, "
        f"project diperbarui {summary.projects_updated}, partner baru {summary.partners_created}, "
        f"baris dilewati {summary.skipped}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
